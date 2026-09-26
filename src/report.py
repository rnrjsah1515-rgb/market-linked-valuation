"""결과물: 차트(PNG), 엑셀 모델(수식 연결), 요약 마크다운."""
from __future__ import annotations

from dataclasses import asdict
from pathlib import Path

import matplotlib
import matplotlib.ticker

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.colors import LinearSegmentedColormap
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

from .model import Assumptions, Wacc

# ── 차트 스타일 (라이트 표면 기준 팔레트) ─────────────────────────
SURFACE, INK, INK2, MUTED, GRID, AXIS = "#fcfcfb", "#0b0b0b", "#52514e", "#898781", "#e1e0d9", "#c3c2b7"
BLUE, ORANGE = "#2a78d6", "#eb6834"
BLUE_ORDINAL = ["#86b6ef", "#3987e5", "#1c5cab", "#0d366b"]   # 헤지비율 0→75% (순서형)
DIVERGING = LinearSegmentedColormap.from_list("div", ["#e34948", "#f0efec", "#2a78d6"])

def _korean_font() -> list[str]:
    from matplotlib import font_manager
    installed = {f.name for f in font_manager.fontManager.ttflist}
    return [f for f in ("Malgun Gothic", "AppleGothic", "NanumGothic", "Noto Sans CJK KR") if f in installed] + ["sans-serif"]


plt.rcParams.update({
    "font.family": _korean_font(),
    "axes.unicode_minus": False, "figure.facecolor": SURFACE, "axes.facecolor": SURFACE,
    "axes.edgecolor": AXIS, "axes.labelcolor": INK2, "xtick.color": MUTED, "ytick.color": MUTED,
    "axes.grid": True, "grid.color": GRID, "grid.linewidth": 0.6, "axes.spines.top": False,
    "axes.spines.right": False, "text.color": INK, "font.size": 10, "axes.titlesize": 12,
    "axes.titleweight": "bold", "axes.titlelocation": "left",
})


COMMA = matplotlib.ticker.FuncFormatter(lambda x, _: f"{x:,.0f}")


def _won(x: float) -> str:
    return f"{x:,.0f}"


def chart_football_field(ranges: pd.DataFrame, base: float, price: float | None, path: Path, demo: bool):
    fig, ax = plt.subplots(figsize=(8, 0.6 * len(ranges) + 1.6))
    y = np.arange(len(ranges))[::-1]
    ax.barh(y, ranges["high"] - ranges["low"], left=ranges["low"], height=0.5, color=BLUE, edgecolor=SURFACE, linewidth=2)
    for yi, (_, r) in zip(y, ranges.iterrows()):
        ax.text(r["low"], yi, f"{_won(r['low'])} ", va="center", ha="right", color=INK2, fontsize=9)
        ax.text(r["high"], yi, f" {_won(r['high'])}", va="center", ha="left", color=INK2, fontsize=9)
    ax.axvline(base, color=INK, linewidth=1.2, linestyle="--")
    box = dict(boxstyle="round,pad=0.25", facecolor=SURFACE, edgecolor="none")
    ax.text(base, len(ranges) - 0.4, f"기준 DCF {_won(base)}원", ha="center", va="bottom", color=INK, fontsize=9, bbox=box, zorder=5)
    if price:
        ax.axvline(price, color=ORANGE, linewidth=1.2)
        ax.text(price, -0.75, f"현재 주가 {_won(price)}원", ha="center", va="top", color=INK2, fontsize=9, bbox=box, zorder=5)
    ax.set_yticks(y, ranges.index)
    ax.xaxis.set_major_formatter(COMMA)
    ax.set_xlabel("주당가치 (원)")
    ax.grid(axis="y", visible=False)
    lo, hi = ranges["low"].min(), ranges["high"].max()
    ax.set_xlim(lo - (hi - lo) * 0.25, hi + (hi - lo) * 0.25)
    ax.set_ylim(-1.2, len(ranges) - 0.1)
    ax.set_title(("[DEMO 가상 데이터] " if demo else "") + "가치평가 범위 (Football Field)")
    fig.text(0.01, 0.01, "주주가치가 음수인 경우 0원으로 표시 (주주 유한책임)", color=MUTED, fontsize=8)
    fig.tight_layout()
    fig.savefig(path, dpi=160)
    plt.close(fig)


def chart_cases(cases: pd.DataFrame, price: float, path: Path, demo: bool):
    """가정 조합별 주당가치와 현재 주가 비교."""
    fig, ax = plt.subplots(figsize=(8.4, 0.62 * len(cases) + 2.0))
    labels = [f"{r['비율 기간']} · {r['할인율']}" for _, r in cases.iterrows()]
    y = np.arange(len(cases))[::-1]
    colors = [BLUE if "채택" in r["할인율"] else "#9ec5f4" for _, r in cases.iterrows()]
    ax.barh(y, cases["주당가치(원)"], height=0.55, color=colors, edgecolor=SURFACE, linewidth=2)
    for yi, v in zip(y, cases["주당가치(원)"]):
        ax.text(v, yi, f" {_won(v)}원 ({v / price - 1:+.0%})", va="center", ha="left", color=INK2, fontsize=9)
    ax.axvline(price, color=ORANGE, linewidth=1.4)
    ax.text(price, len(cases) - 0.45, f"시장가격 {_won(price)}원", ha="center", va="bottom",
            color=ORANGE, fontsize=9, bbox=dict(boxstyle="round,pad=0.25", facecolor=SURFACE, edgecolor="none"), zorder=5)
    ax.set_yticks(y, labels)
    ax.set_xscale("log")
    lo, hi = price * 0.5, max(cases["주당가치(원)"]) * 3
    ticks = [t for t in (5e3, 1e4, 2e4, 5e4, 1e5, 2e5, 5e5, 1e6) if lo <= t <= hi]
    ax.set_xticks(ticks)
    ax.xaxis.set_major_formatter(COMMA)
    ax.xaxis.set_minor_formatter(matplotlib.ticker.NullFormatter())
    ax.set_xlabel("주당가치 (원, 로그 눈금)")
    ax.grid(axis="y", visible=False)
    ax.set_xlim(lo, hi)
    ax.set_ylim(-0.7, len(cases) - 0.05)
    ax.set_title(("[DEMO 가상 데이터] " if demo else "") + "가정 조합별 주당가치")
    fig.tight_layout()
    fig.savefig(path, dpi=160)
    plt.close(fig)


def chart_backtest(df: pd.DataFrame, path: Path):
    """과거 시점별 모델 주당가치 vs 실제 주가."""
    fig, ax = plt.subplots(figsize=(8.6, 4.6))
    x = np.arange(len(df))
    ax.plot(x, df["실제 주가"], marker="o", markersize=8, linewidth=2, color=ORANGE, label="실제 주가")
    ax.plot(x, df["모델 주당가치(8.4%)"], marker="s", markersize=8, linewidth=2, color=BLUE, label="모델 주당가치 (WACC 8.4%)")
    top = df["모델 주당가치(8.4%)"].max()
    for xi, (_, r) in zip(x, df.iterrows()):
        dy = -20 if r["모델 주당가치(8.4%)"] == top else 12     # 최고점 라벨은 제목과 겹치지 않게 아래로
        ax.annotate(f"{r['괴리(8.4%)']:+.0%}", (xi, r["모델 주당가치(8.4%)"]),
                    textcoords="offset points", xytext=(0, dy), ha="center", fontsize=9, color=INK2)
    ax.axhline(0, color=AXIS, linewidth=1)
    ax.set_yscale("symlog", linthresh=10000)
    ax.set_yticks([-40000, -20000, 0, 20000, 40000, 80000])
    ax.yaxis.set_major_formatter(COMMA)
    ax.set_xticks(x, [r["평가일"][:7] + "\n(기준 " + str(r["기준연도"]) + ")" for _, r in df.iterrows()])
    ax.set_ylabel("주당가치 (원, 대칭 로그)")
    ax.legend(frameon=False, fontsize=9, labelcolor=INK2)
    ax.set_title("과거 시점 검증 — 모델 주당가치와 실제 주가")
    fig.tight_layout()
    fig.savefig(path, dpi=160)
    plt.close(fig)


def chart_implied_risk(df: pd.DataFrame, path: Path):
    """시기별 내재 할인율과 관측 가능한 위험지표(무위험이자율·신용스프레드) 비교."""
    fig, ax = plt.subplots(figsize=(8.6, 4.6))
    x = np.arange(len(df))
    ax.bar(x, df["무위험이자율"], width=0.5, color="#cde2fb", label="무위험이자율 (국고채 10년)")
    ax.bar(x, df["신용스프레드(회사채-국고채)"], width=0.5, bottom=df["무위험이자율"],
           color="#86b6ef", label="신용스프레드 (회사채 AA− − 국고채)")
    ax.plot(x, df["내재 WACC(마진수렴)"], marker="o", markersize=8, linewidth=2, color=ORANGE,
            label="주가가 함축하는 할인율 (마진수렴 기준)")
    for xi, v in zip(x, df["내재 WACC(마진수렴)"]):
        ax.annotate(f"{v:.1%}", (xi, v), textcoords="offset points", xytext=(0, 10),
                    ha="center", fontsize=9, color=INK2)
    ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: f"{v:.0%}"))
    ax.set_xticks(x, [r["평가일"][:7] + "\n(기준 " + str(r["기준연도"]) + ")" for _, r in df.iterrows()])
    ax.set_ylabel("연율")
    ax.legend(frameon=False, fontsize=9, labelcolor=INK2, loc="upper left")
    ax.set_title("시기별 내재 할인율 vs 관측된 시장위험")
    fig.tight_layout()
    fig.savefig(path, dpi=160)
    plt.close(fig)


def chart_heatmap(grid: pd.DataFrame, base: float, title: str, path: Path, demo: bool):
    fig, ax = plt.subplots(figsize=(7, 4.6))
    rel = grid.values / base - 1
    lim = max(abs(rel).max(), 1e-9)
    ax.imshow(rel, cmap=DIVERGING, vmin=-lim, vmax=lim, aspect="auto")
    for i in range(grid.shape[0]):
        for j in range(grid.shape[1]):
            ax.text(j, i, f"{grid.values[i, j]:,.0f}\n({rel[i, j]:+.0%})", ha="center", va="center", fontsize=8, color=INK)
    ax.set_xticks(range(grid.shape[1]), grid.columns)
    ax.set_yticks(range(grid.shape[0]), grid.index)
    ax.grid(False)
    for s in ax.spines.values():
        s.set_visible(False)
    ax.set_title(("[DEMO 가상 데이터] " if demo else "") + title)
    fig.tight_layout()
    fig.savefig(path, dpi=160)
    plt.close(fig)


def chart_hedge(ebitda: dict[float, np.ndarray], label: str, base: float, path: Path, demo: bool):
    fig, ax = plt.subplots(figsize=(8, 4.2))
    all_v = np.concatenate(list(ebitda.values()))
    bins = np.linspace(np.percentile(all_v, 0.5), np.percentile(all_v, 99.5), 60)
    for color, (h, v) in zip(BLUE_ORDINAL, ebitda.items()):
        ax.hist(v, bins=bins, histtype="step", linewidth=2, color=color,
                label=f"헤지 {h:.0%}  (표준편차 {_won(v.std())}억원, P5 {_won(np.percentile(v, 5))}억원)")
    ax.axvline(base, color=INK, linewidth=1, linestyle="--")
    ax.xaxis.set_major_formatter(COMMA)
    ax.set_xlabel(f"{label} (억원)")
    ax.set_ylabel("시뮬레이션 경로 수")
    ax.legend(frameon=False, fontsize=9, labelcolor=INK2, loc="upper left")
    ax.set_title(("[DEMO 가상 데이터] " if demo else "") + f"유가 헤지 비율별 {label} 분포 (몬테카를로)")
    fig.tight_layout()
    fig.savefig(path, dpi=160)
    plt.close(fig)


def chart_market(oil: pd.Series, fx: pd.Series, oil_base: float, fx_base: float, path: Path, demo: bool):
    fig, (a1, a2) = plt.subplots(2, 1, figsize=(8, 5.2), sharex=True)
    a1.plot(oil.index, oil.values, color=BLUE, linewidth=2)
    a1.axhline(oil_base, color=INK2, linewidth=1, linestyle="--")
    a1.text(oil.index[0], oil_base, f" 기준 유가 {oil_base:,.1f}", va="bottom", color=INK2, fontsize=9)
    a1.set_title(("[DEMO 가상 데이터] " if demo else "") + "국제유가 (달러/배럴, 월평균)")
    fxm = fx.resample("MS").mean().loc[oil.index[0]:]
    a2.plot(fxm.index, fxm.values, color=BLUE, linewidth=2)
    a2.axhline(fx_base, color=INK2, linewidth=1, linestyle="--")
    a2.text(fxm.index[0], fx_base, f" 기준 환율 {fx_base:,.1f}", va="bottom", color=INK2, fontsize=9)
    a2.set_title("원/달러 환율 (월평균)")
    fig.tight_layout()
    fig.savefig(path, dpi=160)
    plt.close(fig)


# ── 엑셀 모델 ────────────────────────────────────────────────
HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1C5CAB")
INPUT = Font(color="1F4FBF")            # 파란 글씨 = 입력값 (재무모델 관례)
BOLD = Font(bold=True)
THIN = Border(top=Side(style="thin", color="C3C2B7"))
NUM = '#,##0;[Red]-#,##0'
PCT = '0.00%'


def _name(wb: Workbook, name: str, sheet: str, cell: str):
    wb.defined_names[name] = DefinedName(name, attr_text=f"'{sheet}'!${cell[0]}${cell[1:]}")


def _df_sheet(wb: Workbook, title: str, df: pd.DataFrame, index=True, fmt=NUM):
    ws = wb.create_sheet(title)
    df = df.reset_index() if index else df
    for j, col in enumerate(df.columns, 1):
        c = ws.cell(1, j, str(col))
        c.font, c.fill = HEAD, HEAD_FILL
        ws.column_dimensions[get_column_letter(j)].width = max(12, min(40, len(str(col)) * 2 + 2))
    for i, row in enumerate(df.itertuples(index=False), 2):
        for j, v in enumerate(row, 1):
            v = v.item() if hasattr(v, "item") else v
            c = ws.cell(i, j, v)
            if isinstance(v, float):
                c.number_format = fmt
    ws.freeze_panes = "A2"
    return ws


def write_excel(path: Path, a: Assumptions, w: Wacc, hist: pd.DataFrame, mapping: pd.DataFrame,
                sens_wg: pd.DataFrame, sens_of: pd.DataFrame, hedge: pd.DataFrame, summary: dict, demo: bool,
                extra: dict[str, pd.DataFrame] | None = None):
    wb = Workbook()
    ws = wb.active
    ws.title = "가정"
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 60
    ws["A1"] = ("[DEMO 가상 데이터] " if demo else "") + f"{a.company} 밸류에이션 가정"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "파란 글씨 = 입력값, 검은 글씨 = 수식. 입력값을 바꾸면 DCF 시트가 다시 계산됩니다."
    rows = [
        # (이름, 라벨, 값 또는 수식, 서식, 설명)
        ("val_date", "평가기준일", a.valuation_date.to_pydatetime(), "yyyy-mm-dd", "시장 데이터 기준일"),
        ("rf", "무위험이자율 (국고채 10년)", w.risk_free, PCT, "ECOS 817Y002"),
        ("beta", "베타 (재레버리지)", w.relevered_beta, "0.000", f"{w.beta_source}, Blume 조정 후 목표 D/E 로 재레버리지"),
        ("erp", "시장위험프리미엄", w.erp, PCT, "config 가정"),
        ("size_prem", "규모 프리미엄", w.size_premium, PCT, "config 가정"),
        ("ke", "자기자본비용", "=rf+beta*erp+size_prem", PCT, "CAPM"),
        ("kd_pre", "세전 타인자본비용", w.pre_tax_kd, PCT, "ECOS 회사채 AA- 3년 + 신용스프레드"),
        ("tax", "법인세율", w.tax_rate, PCT, "config 가정"),
        ("we", "자기자본 비중", w.equity_weight, PCT, "목표 자본구조"),
        ("capm_wacc", "WACC (CAPM 계산값)", "=we*ke+(1-we)*kd_pre*(1-tax)", PCT, "참고용"),
        ("wacc", "WACC (채택)", w.wacc, PCT,
         "회사 공시값 채택" if w.override is not None else "=CAPM 계산값"),
        ("g", "영구성장률", a.terminal_growth, PCT, ""),
        ("mid", "기중 현금흐름 가정 (1=예)", 1 if a.mid_year else 0, "0", "mid-year convention"),
        ("rev_last", f"기준연도({a.base_year}) 매출 (억원)", a.revenue_last, NUM, "DART"),
        ("fuel_vol_last", "기준연도 유류 소비량 (백만 배럴 상당)", a.fuel_volume_last / 1e6, "#,##0.00",
         "유류비 ÷ (기준연도 평균 유가 × 평균 환율)"),
        ("fx_ref", "기준연도 평균 환율", a.fx_ref, "#,##0.0", "환율 효과의 기준점"),
        ("oil_ref", "기준연도 평균 유가", a.oil_ref, "#,##0.00", "유류할증 전가의 기준점"),
        ("pass_through", "유가 변동 매출 전가율", a.fuel_pass_through, PCT, "유류할증료 구조 (config 가정)"),
        ("s_rev", "매출 달러 연동 비중", a.usd_revenue_share, PCT, "사업보고서 시장위험 주석"),
        ("s_cost", "비유류 비용 달러 연동 비중", a.usd_nonfuel_cost_share, PCT, "사업보고서 시장위험 주석"),
        ("nonfuel_ratio", "비유류 현금비용 / 매출", a.nonfuel_ratio, PCT, "과거 평균"),
        ("da_ratio", "감가상각비 / 매출", a.da_ratio, PCT, "과거 평균"),
        ("capex_ratio", "CAPEX / 매출", a.capex_ratio, PCT, "과거 평균"),
        ("capex_floor", "CAPEX 하한 (감가상각비 대비)", a.capex_floor_to_da, "0.00", "리스 사용권자산 취득 반영"),
        ("nwc_ratio", "영업운전자본 / 매출", a.nwc_ratio, PCT, "과거 평균"),
        ("nwc_last", "기준연도 영업운전자본 (억원)", a.nwc_last, NUM, "DART"),
        ("tcapex_da", "영구기간 CAPEX / 감가상각비", a.terminal_capex_to_da, "0.00", ""),
        ("net_debt", "순차입금 (리스부채 포함, 억원)", a.net_debt, NUM, "DART (재무상태표 공시일 기준)"),
        ("usd_net_debt", "순외화부채 (달러)", a.usd_net_debt, NUM, "사업보고서 시장위험 공시"),
        ("fx_bs", "재무상태표 기준일 환율", a.fx_at_balance_sheet, "#,##0.0", "순차입금이 환산된 시점"),
        ("fx_spot", "평가일 현물환율", a.fx_spot, "#,##0.0", ""),
        ("fx_curve0", "기준 환율 (1차연도)", float(a.fx_curve[0]), "#,##0.0", "환율 시나리오의 기준점"),
        ("nci", "비지배지분 (억원)", a.nci, NUM, "DART"),
        ("shares", "발행주식수 (주)", a.shares, NUM, "주식시세 API 또는 config"),
    ]
    names_row = {}
    for i, (name, label, value, fmt, note) in enumerate(rows, 4):
        names_row[name] = i
        ws.cell(i, 1, label)
        c = ws.cell(i, 2, value)
        c.number_format = fmt
        c.font = BOLD if isinstance(value, str) else INPUT
        ws.cell(i, 3, note).font = Font(color="898781")
        _name(wb, name, "가정", f"B{i}")

    # DCF 시트 (모든 셀 수식)
    d = wb.create_sheet("DCF")
    n = len(a.years)
    cols = [get_column_letter(3 + j) for j in range(n)]
    d.column_dimensions["A"].width = 30
    d.column_dimensions["B"].width = 4
    for c in cols:
        d.column_dimensions[c].width = 14
    d["A1"] = ("[DEMO 가상 데이터] " if demo else "") + f"{a.company} DCF (단위: 억원, 유가 달러/배럴, 환율 원/달러)"
    d["A1"].font = Font(bold=True, size=13)
    d["A2"] = f"유가·환율 커브: {a.curve_source}"
    lines = {}

    def row(r, label, formulas, fmt=NUM, bold=False, font=None):
        d.cell(r, 1, label).font = BOLD if bold else Font()
        for j, f in enumerate(formulas):
            c = d.cell(r, 3 + j, f)
            c.number_format = fmt
            if font:
                c.font = font
            elif bold:
                c.font = BOLD
        lines[label] = r

    row(4, "연도", list(a.years), "0", bold=True)
    for j in range(n):
        d.cell(4, 3 + j).fill = HEAD_FILL
        d.cell(4, 3 + j).font = HEAD
    row(5, "물량 성장률", list(a.growth), PCT, font=INPUT)
    row(6, "누적 성장계수", [f"=1+{c}5" if j == 0 else f"={cols[j-1]}6*(1+{c}5)" for j, c in enumerate(cols)], "0.0000")
    row(7, "유가", list(a.oil_curve), "#,##0.00", font=INPUT)
    row(8, "환율", list(a.fx_curve), "#,##0.0", font=INPUT)
    row(9, "매출 (유류할증 전가 포함)", [f"=rev_last*{c}6*(1+s_rev*({c}8/fx_ref-1))+pass_through*{c}10*({c}7-oil_ref)*{c}8/100"
                                   for c in cols], bold=True)
    row(10, "유류 소비량 (백만 배럴 상당)", [f"=fuel_vol_last*{c}6" for c in cols], "#,##0.00")
    row(11, "유류비", [f"={c}10*{c}7*{c}8/100" for c in cols])
    row(12, "비유류 현금비용", [f"=nonfuel_ratio*rev_last*{c}6*(1+s_cost*({c}8/fx_ref-1))" for c in cols])
    row(13, "EBITDA", [f"={c}9-{c}11-{c}12" for c in cols], bold=True)
    row(14, "감가상각비", [f"=da_ratio*{c}9" for c in cols])
    row(15, "EBIT", [f"={c}13-{c}14" for c in cols])
    row(16, "법인세", [f"=tax*MAX({c}15,0)" for c in cols])
    row(17, "NOPAT", [f"={c}15-{c}16" for c in cols])
    row(18, "CAPEX", [f"=MAX(capex_ratio*{c}9,capex_floor*{c}14)" for c in cols])
    row(19, "영업운전자본", [f"=nwc_ratio*{c}9" for c in cols])
    row(20, "운전자본 증감", [f"={c}19-nwc_last" if j == 0 else f"={c}19-{cols[j-1]}19" for j, c in enumerate(cols)])
    row(21, "FCFF", [f"={c}17+{c}14-{c}18-{c}20" for c in cols], bold=True)
    row(22, "반영비율 (평가일 이후 기간)", [f"=MIN(MAX((DATE({c}4,12,31)-val_date)/365.25,0),1)" for c in cols], "0.000")
    row(23, "할인기간 (년)", [f"=(DATE({c}4,12,31)-val_date)/365.25-IF(mid=1,{c}22/2,0)" for c in cols], "0.000")
    row(24, "할인계수", [f"=1/(1+wacc)^{c}23" for c in cols], "0.0000")
    row(25, "FCFF 현재가치", [f"={c}21*{c}22*{c}24" for c in cols], bold=True)
    for c in cols:
        d[f"{c}21"].border = THIN

    L = cols[-1]
    tv_rows = [
        ("영구기간 매출", f"={L}9*(1+g)", NUM),
        ("영구기간 EBITDA", f"={L}13*(1+g)", NUM),
        ("영구기간 감가상각비", "=da_ratio*C28", NUM),
        ("영구기간 EBIT", "=C29-C30", NUM),
        ("영구기간 NOPAT", "=C31-tax*MAX(C31,0)", NUM),
        ("영구기간 CAPEX", "=tcapex_da*C30", NUM),
        ("영구기간 운전자본 증감", f"=nwc_ratio*(C28-{L}9)", NUM),
        ("영구기간 FCFF", "=C32+C30-C33-C34", NUM),
        ("잔존가치 (영구성장)", "=C35/(wacc-g)", NUM),
        ("잔존가치 할인기간", f"=(DATE({L}4,12,31)-val_date)/365.25-IF(mid=1,0.5,0)", "0.000"),
        ("잔존가치 현재가치", "=C36/(1+wacc)^C37", NUM),
        ("예측기간 FCFF 현재가치 합계", f"=SUM(C25:{L}25)", NUM),
        ("기업가치 (EV)", "=C38+C39", NUM),
        ("(-) 순차입금 (공시)", "=net_debt", NUM),
        ("(-) 환율 재평가 (순외화부채)", "=usd_net_debt*(fx_spot*C8/fx_curve0-fx_bs)/100000000", NUM),
        ("(-) 비지배지분", "=nci", NUM),
        ("주주가치", "=C40-C41-C42-C43", NUM),
        ("주당가치 (원)", "=C44*100000000/shares", NUM),
        ("잔존가치 비중", "=C38/C40", PCT),
    ]
    for i, (label, f, fmt) in enumerate(tv_rows, 28):
        d.cell(i, 1, label)
        c = d.cell(i, 3, f)
        c.number_format = fmt
        if label in ("기업가치 (EV)", "주주가치", "주당가치 (원)"):
            d.cell(i, 1).font = c.font = BOLD
    d.freeze_panes = "C5"

    _df_sheet(wb, "과거재무", hist.round(1))
    _df_sheet(wb, "WACC_상세", pd.DataFrame({"항목": list(asdict(w).keys()), "값": list(asdict(w).values())}), index=False, fmt="0.0000")
    if extra is not None:
        for title, df in extra.items():
            _df_sheet(wb, title, df, index=False, fmt='#,##0.0')
    _df_sheet(wb, "민감도_WACC_g", sens_wg)
    _df_sheet(wb, "민감도_유가_환율", sens_of)
    _df_sheet(wb, "헤지분석", hedge, index=False, fmt="#,##0.00")
    if not mapping.empty:
        _df_sheet(wb, "DART_매핑로그", mapping, index=False, fmt="#,##0.0")
    s = wb.create_sheet("파이썬_검증값", 1)
    s["A1"], s["B1"] = "항목", "파이썬 계산값"
    for i, (k, v) in enumerate(summary.items(), 2):
        s.cell(i, 1, k)
        s.cell(i, 2, v).number_format = NUM if abs(v) > 10 else "0.0000"
    s.column_dimensions["A"].width = 28
    s.column_dimensions["B"].width = 18
    for sheet in wb.worksheets:
        for r in sheet.iter_rows():
            for c in r:
                if c.column > 1:
                    c.alignment = Alignment(horizontal="right")
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return {"per_share_cell": "DCF!C45", "ev_cell": "DCF!C40", "wacc_cell": f"가정!B{names_row['wacc']}"}


# ── 요약 마크다운 ─────────────────────────────────────────────
def _md_table(df: pd.DataFrame, fmt="{:,.0f}") -> str:
    df = df.copy()
    head = "| " + " | ".join([df.index.name or ""] + [str(c) for c in df.columns]) + " |"
    sep = "|" + "---|" * (len(df.columns) + 1)
    body = []
    for idx, r in df.iterrows():
        cells = [fmt.format(v) if isinstance(v, (int, float, np.floating)) else str(v) for v in r.values]
        body.append("| " + " | ".join([str(idx)] + cells) + " |")
    return "\n".join([head, sep, *body])


def write_summary(path: Path, a: Assumptions, w: Wacc, base: dict, ranges: pd.DataFrame,
                  sens_of: pd.DataFrame, hedge: pd.DataFrame, price: float | None, demo: bool, excel_check: str,
                  implied: pd.DataFrame | None = None, cases: pd.DataFrame | None = None):
    proj = base["projection"]
    fc = pd.DataFrame({k: proj[k][0] for k in ("revenue", "fuel_cost", "ebitda", "fcff")}, index=a.years).T
    fc.index = ["매출", "유류비", "EBITDA", "FCFF"]
    fc.index.name = "억원"
    warn = "> **DEMO 모드: 모든 수치는 형식 확인용 가상 데이터입니다. 실제 기업 평가가 아닙니다.**\n\n" if demo else ""
    upside = f"{base['per_share'][0] / price - 1:+.1%}" if price else "-"
    hedge_md = hedge.set_index("헤지비율").copy()
    for col in [c for c in hedge_md.columns if "확률" in c]:
        hedge_md[col] = hedge_md[col].map(lambda v: f"{v:.1%}")
    implied_md = _md_table(implied.set_index("항목"), "{}") if implied is not None else "(생략)"
    cases_md = _md_table(cases.set_index("비율 기간"), "{:,.0f}") if cases is not None else "(생략)"
    wacc_note = "회사 공시값 채택" if w.override is not None else "CAPM"
    text = f"""# {a.company} 밸류에이션 요약
{warn}
- 평가기준일: {a.valuation_date.date()} / 기준 재무연도: {a.base_year}
- 유가·환율 커브: {a.curve_source}
- 엑셀 수식 검증: {excel_check}

## 결과

| 항목 | 값 |
|---|---|
| 기업가치 (EV) | {base['ev'][0]:,.0f} 억원 |
| 순차입금 (리스부채 포함) | {a.net_debt:,.0f} 억원 |
| 주주가치 | {base['equity'][0]:,.0f} 억원 |
| **주당가치** | **{base['per_share'][0]:,.0f} 원** |
| 현재 주가 | {f'{price:,.0f} 원' if price else '-'} |
| 괴리율 | {upside} |
| 잔존가치 비중 | {base['pv_tv'][0] / base['ev'][0]:.1%} |

## 시장가격 역산 — 지금 주가를 설명하려면 무엇을 가정해야 하는가

{implied_md}

다른 가정을 그대로 두고 한 항목만 움직였을 때, 시가총액과 같아지는 값입니다.

## 가정 조합별 주당가치

{cases_md}

![cases](charts/cases.png)

## 할인율

| 항목 | 값 |
|---|---|
| **채택 할인율** | **{w.wacc:.2%}** ({wacc_note}) |
| CAPM 계산값 | {w.capm_wacc:.2%} |
| 무위험이자율 (국고채 10년) | {w.risk_free:.2%} |
| 베타 (원시 → Blume → 재레버리지) | {w.raw_beta:.3f} → {w.adjusted_beta:.3f} → {w.relevered_beta:.3f} |
| 베타 추정 방법 | {w.beta_source} |
| 시장위험프리미엄 | {w.erp:.2%} |
| 자기자본비용 | {w.cost_of_equity:.2%} |
| 세후 타인자본비용 | {w.after_tax_kd:.2%} |
| 자기자본 비중 | {w.equity_weight:.1%} |
| **WACC** | **{w.wacc:.2%}** |

## 추정 손익 (기준 커브)

{_md_table(fc)}

## 가치평가 범위 (주당가치, 원)

{_md_table(ranges[["low", "high"]].rename(columns={"low": "하단", "high": "상단"}))}

## 유가 × 환율 민감도 (주당가치, 원)

{_md_table(sens_of)}

## 유가 헤지 비율별 분포

{_md_table(hedge_md)}

헤지 비용이 없고 선물가격이 기대 현물가격과 같다고 가정하면, 헤지로 **평균 가치는 거의 변하지 않고 분포의 폭(하방 위험)만 줄어듭니다.**
헤지가 기업가치를 높이는 경로는 재무적 곤경 비용, 차입 여력, 세금 비선형성 같은 시장 마찰에서 나옵니다.

## 차트

![football](charts/football_field.png)
![oil-fx](charts/sensitivity_oil_fx.png)
![hedge](charts/hedge_ebitda.png)
![market](charts/market_history.png)
"""
    path.write_text(text, encoding="utf-8")
